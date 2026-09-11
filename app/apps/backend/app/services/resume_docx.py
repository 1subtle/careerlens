"""Editable Word paragraphs, runs and lists from the same saved print settings."""

from html.parser import HTMLParser
from io import BytesIO
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.schemas.models import ResumeData, normalize_resume_data
from app.schemas.template_settings import TemplateSettings
from app.services.resume_photo import decode_resume_photo

# Same CSS px/rem and compact mappings as frontend/lib/types/template-settings.ts.
BASE_PT = (8.25, 9, 10.5, 11.25, 12)
SECTION_PT = (4.5, 7.5, 12, 15, 18)
ITEM_PT = (1.5, 3, 6, 9, 12)
LINE_HEIGHT = (1.15, 1.25, 1.35, 1.45, 1.55)
HEADER_SCALE = (1.5, 1.75, 2, 2.25, 2.5)
SECTION_SCALE = (1, 1.1, 1.2, 1.3, 1.4)
ACCENT = {"blue": "1D4ED8", "green": "15803D", "orange": "EA580C", "red": "DC2626"}
FONTS = {
    "serif": ("Georgia", "SimSun"),
    "sans-serif": ("Arial", "Microsoft YaHei"),
    "mono": ("Consolas", "Microsoft YaHei"),
}
ZH_HEADINGS = {
    "summary": ("Summary", "个人简介"),
    "workExperience": ("Experience", "工作经历"),
    "education": ("Education", "教育背景"),
    "personalProjects": ("Projects", "项目经历"),
    "additional": ("Skills & Awards", "技能与荣誉"),
}


class RichText(HTMLParser):
    """Keep text and basic emphasis as editable runs; discard active HTML."""

    def __init__(self, value: str):
        super().__init__(convert_charrefs=True)
        self.blocks: list[list[tuple[str, dict[str, bool]]]] = [[]]
        self.marks = {"bold": 0, "italic": 0, "underline": 0}
        self.hidden = 0
        self.feed(value)
        self.close()

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in ("script", "style", "iframe", "object", "svg"):
            self.hidden += 1
        if self.hidden:
            return
        mark = {
            "b": "bold",
            "strong": "bold",
            "i": "italic",
            "em": "italic",
            "u": "underline",
        }.get(tag)
        if mark:
            self.marks[mark] += 1
        if tag in ("p", "div", "li") and self.blocks[-1]:
            self.blocks.append([])
        if tag == "br":
            self.handle_data("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "iframe", "object", "svg"):
            self.hidden = max(0, self.hidden - 1)
            return
        if self.hidden:
            return
        mark = {
            "b": "bold",
            "strong": "bold",
            "i": "italic",
            "em": "italic",
            "u": "underline",
        }.get(tag)
        if mark:
            self.marks[mark] = max(0, self.marks[mark] - 1)
        if tag in ("p", "div", "li") and self.blocks[-1]:
            self.blocks.append([])

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.blocks[-1].append(
                (data, {key: bool(value) for key, value in self.marks.items()})
            )

    def text_blocks(self) -> list[list[tuple[str, dict[str, bool]]]]:
        return [
            block for block in self.blocks if any(text.strip() for text, _ in block)
        ]


def _font_properties(properties: Any, family: str) -> None:
    latin, chinese = FONTS[family]
    fonts = properties.get_or_add_rFonts()
    # Theme font attributes take precedence over explicit names in Word/LibreOffice.
    fonts.attrib.clear()
    for key, value in (
        ("ascii", latin),
        ("hAnsi", latin),
        ("cs", latin),
        ("eastAsia", chinese),
        ("hint", "eastAsia"),
    ):
        fonts.set(qn("w:" + key), value)
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "zh-CN")


def _font(
    style: Any, family: str, size: float, *, bold: bool = False, color: str = "262626"
) -> None:
    style.font.name = FONTS[family][0]
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    properties = style.element.get_or_add_rPr()
    _font_properties(properties, family)
    complex_size = properties.find(qn("w:szCs"))
    if complex_size is not None:
        properties.remove(complex_size)


def _plain(value: str | None) -> str:
    return "\n".join(
        "".join(text for text, _ in block)
        for block in RichText(value or "").text_blocks()
    ).strip()


def _no_grid(properties: Any) -> None:
    snap = properties.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        properties.append(snap)
    snap.set(qn("w:val"), "0")


def _line_height(paragraph_format: Any, points: float) -> None:
    # Match CSS line boxes without CJK fallback fonts inflating their height.
    # Inline photos use automatic spacing separately so their height is preserved.
    paragraph_format.line_spacing = Pt(points)
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def render_resume_docx(
    data: ResumeData | dict,
    settings: TemplateSettings | dict | None = None,
    *,
    lang: str = "zh",
) -> bytes:
    resume = data if isinstance(data, ResumeData) else ResumeData.model_validate(data)
    options = (
        settings
        if isinstance(settings, TemplateSettings)
        else TemplateSettings.model_validate(settings or {})
    )
    raw = normalize_resume_data(resume.model_dump(mode="json"))
    document = Document()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        document.styles.element.insert(0, defaults)
    run_default = defaults.find(qn("w:rPrDefault"))
    if run_default is None:
        run_default = OxmlElement("w:rPrDefault")
        defaults.append(run_default)
    run_properties = run_default.find(qn("w:rPr"))
    if run_properties is None:
        run_properties = OxmlElement("w:rPr")
        run_default.append(run_properties)
    _font_properties(run_properties, options.fontSize.bodyFont)
    document.core_properties.title = resume.personalInfo.name or ""
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    section = document.sections[0]
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)
    section.page_width, section.page_height = (
        (Mm(210), Mm(297)) if options.pageSize == "A4" else (Mm(215.9), Mm(279.4))
    )
    for edge, value in options.margins.model_dump().items():
        setattr(section, edge + "_margin", Mm(value))
    base = BASE_PT[options.fontSize.base - 1]
    gap = ITEM_PT[options.spacing.item - 1] * (0.6 if options.compactMode else 1)
    section_gap = SECTION_PT[options.spacing.section - 1] * (
        0.6 if options.compactMode else 1
    )
    line_height = LINE_HEIGHT[options.spacing.lineHeight - 1] * (
        0.92 if options.compactMode else 1
    )
    paragraph_default = defaults.find(qn("w:pPrDefault"))
    if paragraph_default is None:
        paragraph_default = OxmlElement("w:pPrDefault")
        defaults.append(paragraph_default)
    paragraph_properties = paragraph_default.find(qn("w:pPr"))
    if paragraph_properties is None:
        paragraph_properties = OxmlElement("w:pPr")
        paragraph_default.append(paragraph_properties)
    _no_grid(paragraph_properties)
    default_spacing = paragraph_properties.find(qn("w:spacing"))
    if default_spacing is None:
        default_spacing = OxmlElement("w:spacing")
        paragraph_properties.append(default_spacing)
    default_spacing.attrib.clear()
    for key, value in (
        ("before", "0"),
        ("after", "0"),
        ("line", str(round(base * line_height * 20))),
        ("lineRule", "exact"),
    ):
        default_spacing.set(qn("w:" + key), value)
    width = section.page_width - section.left_margin - section.right_margin
    normal = document.styles["Normal"]
    _font(normal, options.fontSize.bodyFont, base)
    _line_height(normal.paragraph_format, base * line_height)
    normal.paragraph_format.space_before = Pt(0)
    _no_grid(normal.element.get_or_add_pPr())
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True
    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        _font(style, options.fontSize.bodyFont, base * 0.92)
        _line_height(style.paragraph_format, base * 0.92 * line_height)
        _no_grid(style.element.get_or_add_pPr())
        style.paragraph_format.space_after = Pt(gap * 0.75)
        style.paragraph_format.left_indent = Mm(4)
        style.paragraph_format.first_line_indent = Mm(-2.5)
    title = document.styles["Title"]
    _font(
        title,
        options.fontSize.headerFont,
        base * HEADER_SCALE[options.fontSize.headerScale - 1],
        bold=options.template != "clean",
        color="111111",
    )
    title.paragraph_format.space_after = Pt(gap * 0.6)
    title.paragraph_format.keep_with_next = True
    _line_height(
        title.paragraph_format,
        base * HEADER_SCALE[options.fontSize.headerScale - 1] * line_height,
    )
    _no_grid(title.element.get_or_add_pPr())
    title.font.underline = False
    for border in list(title.element.get_or_add_pPr().findall(qn("w:pBdr"))):
        title.element.get_or_add_pPr().remove(border)
    heading = document.styles["Heading 1"]
    color = (
        ACCENT[options.accentColor]
        if options.template.startswith("modern") or options.template == "vivid"
        else ("666666" if options.template == "clean" else "111111")
    )
    _font(
        heading,
        options.fontSize.headerFont,
        base
        * SECTION_SCALE[options.fontSize.headerScale - 1]
        * (1.15 if options.template == "clean" else 1),
        bold=True,
        color=color,
    )
    heading.paragraph_format.space_before = Pt(section_gap)
    heading.paragraph_format.space_after = Pt(gap)
    heading.paragraph_format.keep_with_next = True
    _line_height(
        heading.paragraph_format,
        base
        * SECTION_SCALE[options.fontSize.headerScale - 1]
        * (1.15 if options.template == "clean" else 1)
        * line_height,
    )
    _no_grid(heading.element.get_or_add_pPr())
    borders = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "12" if options.template.startswith("modern") else "6")
    border.set(
        qn("w:color"), color if options.template.startswith("modern") else "D1D5DB"
    )
    border.set(qn("w:space"), "1")
    borders.append(border)
    heading.element.get_or_add_pPr().append(borders)

    def body(value: str, *, bullet: bool = False, small: bool = False) -> None:
        for block in RichText(value).text_blocks():
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet else "Normal"
            )
            paragraph.paragraph_format.space_after = Pt(gap * 0.75)
            _line_height(
                paragraph.paragraph_format, base * (0.92 if small else 1) * line_height
            )
            for text, marks in block:
                run = paragraph.add_run(text)
                run.bold, run.italic, run.underline = (
                    marks["bold"],
                    marks["italic"],
                    marks["underline"],
                )
                if small:
                    run.font.size = Pt(base * 0.92)

    def entry(item: dict) -> None:
        primary = item.get("name") or item.get("title") or item.get("institution") or ""
        secondary = (
            item.get("company")
            or item.get("degree")
            or item.get("role")
            or item.get("subtitle")
            or ""
        )
        years = _plain(item.get("years"))
        location = _plain(item.get("location"))
        if primary or years:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(gap * 0.6)
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                width, WD_TAB_ALIGNMENT.RIGHT
            )
            paragraph.add_run(_plain(primary)).bold = True
            if years:
                run = paragraph.add_run("\t" + years)
                run.font.size = Pt(base * 0.82)
                run.font.color.rgb = RGBColor.from_string("666666")
        if secondary or location:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = bool(item.get("description"))
            paragraph.paragraph_format.space_after = Pt(gap)
            _line_height(paragraph.paragraph_format, base * 0.95 * line_height)
            paragraph.add_run(
                " · ".join(value for value in (_plain(secondary), location) if value)
            ).font.size = Pt(base * 0.95)
        description = item.get("description") or []
        if isinstance(description, str):
            body(description, small=True)
        else:
            styles = item.get("descriptionStyles") or []
            for index, text in enumerate(description):
                body(
                    text,
                    bullet=index >= len(styles) or styles[index] != "plain",
                    small=True,
                )
        links = [item.get(key) for key in ("website", "github") if item.get(key)]
        if links:
            body(" | ".join(links), small=True)
        if document.paragraphs:
            document.paragraphs[-1].paragraph_format.space_after = Pt(gap)

    personal = resume.personalInfo
    if personal.photo:
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width, table.columns[1].width = width - Mm(35), Mm(35)
        table.cell(0, 0).width, table.cell(0, 1).width = width - Mm(35), Mm(35)
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                value = OxmlElement("w:" + edge)
                value.set(qn("w:w"), "0")
                value.set(qn("w:type"), "dxa")
                margins.append(value)
            cell._tc.get_or_add_tcPr().append(margins)
        header = table.cell(0, 0)
        image_bytes = decode_resume_photo(personal.photo)
        photo = Image.from_blob(image_bytes)
        scale = min(30 / photo.px_width, 40 / photo.px_height)
        picture_p = table.cell(0, 1).paragraphs[0]
        picture_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        picture_p.paragraph_format.line_spacing = 1.0
        picture_p.add_run().add_picture(
            BytesIO(image_bytes),
            width=Mm(photo.px_width * scale),
            height=Mm(photo.px_height * scale),
        )
        name_p = header.paragraphs[0]
        name_p.style = title
        name_p.add_run(_plain(personal.name))
    else:
        header = document
        name_p = document.add_paragraph(_plain(personal.name), style="Title")
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if personal.title:
        p = header.add_paragraph(_plain(personal.title))
        p.paragraph_format.space_after = Pt(gap)
        _line_height(p.paragraph_format, base * 1.15 * line_height)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(base * 1.15)
    if any(item['key'] == 'education' and item['isVisible'] for item in raw['sectionMeta']):
        def education_rank(item):
            years = item.years or ''
            return 9999 if re.search(r'至今|在读|present|current', years, re.I) else max([0] + [int(year) for year in re.findall(r'(?:19|20)\d{2}', years)])
        education = sorted(resume.education, key=education_rank, reverse=True)
        if education:
            school = education[0]
            labels = ('院校', '学历 / 专业') if lang.startswith('zh') else ('School', 'Degree')
            text = '  |  '.join(f'{label}：{_plain(value)}' for label, value in zip(labels, [school.institution, school.degree]) if value)
            if text:
                p = header.add_paragraph(text)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_labels = [('email', '邮箱', 'Email'), ('phone', '电话', 'Phone'), ('location', '所在地', 'Location'), ('website', '个人网站', 'Website'), ('linkedin', 'LinkedIn', 'LinkedIn'), ('github', 'GitHub', 'GitHub')]
    contacts = [f'{zh if lang.startswith("zh") else en}：{_plain(getattr(personal, key))}'
                for key, zh, en in contact_labels if getattr(personal, key)]
    if contacts:
        p = header.add_paragraph(" | ".join(contacts))
        p.paragraph_format.space_after = Pt(gap * 1.25)
        _line_height(p.paragraph_format, base * 0.9 * line_height)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(base * 0.9)

    for meta in sorted(raw["sectionMeta"], key=lambda item: item["order"]):
        key = meta["key"]
        if not meta["isVisible"] or key == "personalInfo":
            continue
        value = raw.get(key) if meta["isDefault"] else raw["customSections"].get(key)
        if not value or (
            isinstance(value, dict)
            and not any(
                value.get(field)
                for field in (
                    "text",
                    "items",
                    "strings",
                    "technicalSkills",
                    "languages",
                    "certificationsTraining",
                    "awards",
                )
            )
        ):
            continue
        display = meta["displayName"]
        if (
            lang.startswith("zh")
            and key in ZH_HEADINGS
            and display == ZH_HEADINGS[key][0]
        ):
            display = ZH_HEADINGS[key][1]
        document.add_paragraph(display, style="Heading 1")
        if key == "additional":
            for field, label in (
                ("technicalSkills", "技能"),
                ("languages", "语言"),
                ("certificationsTraining", "证书与培训"),
                ("awards", "荣誉"),
            ):
                if value.get(field):
                    body(label + "：" + "、".join(value[field]), small=True)
        elif isinstance(value, str):
            body(value)
        elif isinstance(value, list):
            for item in value:
                entry(item)
        elif meta["sectionType"] == "text":
            body(value.get("text") or "")
        elif meta["sectionType"] == "stringList":
            body("、".join(value.get("strings") or []), small=True)
        else:
            for item in value.get("items") or []:
                entry(item)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
