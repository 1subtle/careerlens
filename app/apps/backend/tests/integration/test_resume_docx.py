"""Saved print settings and editable Word exports use isolated test databases."""

import base64
import copy
import struct
import zlib
from io import BytesIO
from zipfile import ZipFile

import pytest
from app.database import Database
from app.main import app
from app.schemas.models import PersonalInfo, ResumeData
from app.schemas.template_settings import TemplateSettings
from app.services.resume_docx import render_resume_docx
from docx import Document
from httpx import ASGITransport, AsyncClient
from pydantic import ValidationError


def png_data(width=3, height=4):
    def chunk(kind, content):
        return (
            struct.pack(">I", len(content))
            + kind
            + content
            + struct.pack(">I", zlib.crc32(kind + content))
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress((b"\x00" + b"\x80\x80\x80" * width) * height))
        + chunk(b"IEND", b"")
    )


def photo_url(data=None, mime="png"):
    return (
        f"data:image/{mime};base64,"
        + base64.b64encode(data if data is not None else png_data()).decode()
    )


@pytest.fixture
async def client():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as value:
        yield value


def content():
    return {
        "personalInfo": {
            "name": "林同学",
            "title": "产品助理",
            "email": "lin@example.test",
            "photo": photo_url(),
        },
        "summary": "<p>具有<strong>访谈整理</strong>经验。</p><p>关注使用者反馈。</p>",
        "workExperience": [
            {"title": "隐藏经历", "description": ["不应导出的隐藏内容"]}
        ],
        "education": [
            {
                "institution": "示例大学",
                "degree": "信息管理本科",
                "years": "2022–2026",
                "description": "主修管理学",
            }
        ],
        "personalProjects": [
            {
                "name": "用户调研",
                "role": "项目成员",
                "years": "2025",
                "description": [
                    "访谈用户，整理问题清单。",
                    "<em>参与讨论</em>并形成记录。",
                ],
                "descriptionStyles": ["bullet", "plain"],
            }
        ],
        "sectionMeta": [
            {
                "id": "summary",
                "key": "summary",
                "displayName": "个人简介",
                "sectionType": "text",
                "order": 3,
            },
            {
                "id": "work",
                "key": "workExperience",
                "displayName": "工作经历",
                "sectionType": "itemList",
                "order": 1,
                "isVisible": False,
            },
            {
                "id": "edu",
                "key": "education",
                "displayName": "教育背景",
                "sectionType": "itemList",
                "order": 0,
            },
            {
                "id": "projects",
                "key": "personalProjects",
                "displayName": "项目经历",
                "sectionType": "itemList",
                "order": 2,
            },
            {
                "id": "custom",
                "key": "custom",
                "displayName": "志愿服务",
                "sectionType": "text",
                "isDefault": False,
                "order": 4,
            },
        ],
        "customSections": {
            "custom": {
                "sectionType": "text",
                "text": "整理活动记录。<script>doNotExport()</script><style>bad-style</style>",
            }
        },
    }


async def test_save_export_native_word_keeps_order_hidden_sections_and_photo(client):
    settings = TemplateSettings(template="modern", accentColor="green").model_dump()
    saved = await client.post(
        "/api/v1/career/resumes",
        json={
            "title": "林同学 · 产品助理",
            "data": content(),
            "template_settings": settings,
        },
    )
    assert saved.status_code == 200, saved.text
    resume = saved.json()
    assert resume["template_settings"] == settings
    response = await client.get(f"/api/v1/resumes/{resume['id']}/docx")
    assert response.status_code == 200, response.text
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    document = Document(BytesIO(response.content))
    assert len(document.inline_shapes) == 1
    assert len(document.tables) == 1
    header = "\n".join(p.text for p in document.tables[0].cell(0, 0).paragraphs)
    assert "林同学" in header and "lin@example.test" in header
    texts = [p.text for p in document.paragraphs]
    joined = "\n".join(texts)
    assert (
        joined.index("教育背景")
        < joined.index("项目经历")
        < joined.index("个人简介")
        < joined.index("志愿服务")
    )
    assert "不应导出" not in joined and "隐藏经历" not in joined
    assert "doNotExport" not in joined and "bad-style" not in joined
    assert any(
        p.style.name == "List Bullet" and "访谈用户" in p.text
        for p in document.paragraphs
    )
    assert any(
        p.style.name == "Normal" and "参与讨论" in p.text for p in document.paragraphs
    )
    assert any(
        run.bold and "访谈整理" in run.text
        for p in document.paragraphs
        for run in p.runs
    )
    assert any(
        run.italic and "参与讨论" in run.text
        for p in document.paragraphs
        for run in p.runs
    )
    assert document.sections[0].page_width.mm == pytest.approx(210, abs=0.1)
    assert document.sections[0].left_margin.mm == pytest.approx(14, abs=0.1)
    assert document.styles["Normal"].paragraph_format.line_spacing.pt == pytest.approx(
        11.25 * 1.45, abs=0.03
    )
    assert str(document.styles["Heading 1"].font.color.rgb) == "15803D"
    with ZipFile(BytesIO(response.content)) as package:
        xml = package.read("word/document.xml").decode()
        assert "<w:t>" in xml
        assert "<w:drawing>" in xml
        assert not any(
            name.startswith("word/embeddings") for name in package.namelist()
        )


async def test_settings_update_is_atomic_optional_and_does_not_change_content_hash(
    client,
):
    created = (
        await client.post(
            "/api/v1/career/resumes", json={"data": content(), "title": "初版"}
        )
    ).json()
    settings = TemplateSettings(
        template="clean", compactMode=True, margins={"top": 20}
    ).model_dump()
    saved = await client.put(
        f"/api/v1/career/resumes/{created['id']}",
        json={
            "data": created["data"],
            "title": created["title"],
            "expected_hash": created["hash"],
            "template_settings": settings,
        },
    )
    assert saved.status_code == 200
    assert saved.json()["hash"] == created["hash"]
    kept = await client.put(
        f"/api/v1/career/resumes/{created['id']}",
        json={
            "data": created["data"],
            "title": created["title"],
            "expected_hash": created["hash"],
        },
    )
    assert kept.json()["template_settings"] == settings
    bad = await client.put(
        f"/api/v1/career/resumes/{created['id']}",
        json={
            "data": {**created["data"], "summary": "不可半保存"},
            "title": "错误排版",
            "template_settings": {"margins": {"top": 100}},
        },
    )
    assert bad.status_code == 422
    state = (await client.get("/api/v1/career/state")).json()["resumes"][0]
    assert state["data"] == created["data"] and state["template_settings"] == settings
    cleared = await client.put(
        f"/api/v1/career/resumes/{created['id']}",
        json={
            "data": created["data"],
            "title": created["title"],
            "template_settings": None,
        },
    )
    assert cleared.json()["template_settings"] is None


async def test_rewrite_variant_inherits_saved_print_settings(client):
    settings = TemplateSettings(template="clean").model_dump()
    data = content()
    data["additional"] = {"technicalSkills": ["SQL"]}
    resume = (
        await client.post(
            "/api/v1/career/resumes",
            json={"title": "原简历", "data": data, "template_settings": settings},
        )
    ).json()
    job = (
        await client.post(
            "/api/v1/career/jobs",
            json={"title": "产品助理", "text": "使用SQL整理用户需求。"},
        )
    ).json()
    match = (
        await client.post(
            "/api/v1/career/matches",
            json={"resume_id": resume["id"], "job_id": job["job_id"]},
        )
    ).json()
    draft = (
        await client.post(
            "/api/v1/career/rewrites",
            json={"match_id": match["id"], "section_id": "personalProjects:0:0"},
        )
    ).json()
    applied = await client.post(
        f"/api/v1/career/rewrites/{draft['id']}/apply", json={"confirmed": True}
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["resume"]["template_settings"] == settings
    assert (
        applied.json()["resume"]["data"]["personalInfo"]["photo"]
        == data["personalInfo"]["photo"]
    )


async def test_facade_clone_inherits_print_settings(isolated_backend_state):
    db = isolated_backend_state
    settings = TemplateSettings(template="modern").model_dump()
    source = await db.create_resume(
        "original", processed_data=content(), template_settings=settings
    )
    clone = await db.create_resume(
        "clone", parent_id=source["resume_id"], processed_data=content()
    )
    assert clone["template_settings"] == settings


async def test_template_settings_additive_migration_preserves_existing_data(tmp_path):
    path = tmp_path / "old.sqlite"
    db = Database(path)
    saved = await db.create_resume("preserve me", title="旧简历")
    with db._sync_engine.begin() as connection:
        connection.exec_driver_sql("ALTER TABLE resumes DROP COLUMN template_settings")
        connection.exec_driver_sql("DROP TABLE schema_migrations")
        connection.exec_driver_sql("DROP TABLE direction_history")
        connection.exec_driver_sql("DROP TABLE market_history")
    await db.close()
    for _ in range(2):
        reopened = Database(path)
        actual = await reopened.get_resume(saved["resume_id"])
        assert actual["content"] == "preserve me"
        assert actual["template_settings"] is None
        await reopened.close()


@pytest.mark.parametrize("value", [None, ""])
def test_absent_photo_keeps_legacy_serialized_data(value):
    assert "photo" not in PersonalInfo(photo=value).model_dump()
    assert "photo" not in ResumeData().model_dump()["personalInfo"]


@pytest.mark.parametrize(
    "value",
    [
        "https://example.test/photo.png",
        "data:image/svg+xml;base64,PHN2Zz4=",
        "data:image/png;base64,invalid",
        photo_url(mime="jpeg"),
        photo_url(png_data(1601, 1)),
        photo_url(png_data() + b"x" * (1024 * 1024)),
    ],
)
def test_photo_rejects_external_fake_oversized_and_invalid_images(value):
    with pytest.raises(ValidationError):
        PersonalInfo(photo=value)


async def test_export_missing_or_unparsed_resume_is_explicit(
    client, isolated_backend_state
):
    assert (await client.get("/api/v1/resumes/missing/docx")).status_code == 404
    pending = await isolated_backend_state.create_resume("raw only")
    assert (
        await client.get(f"/api/v1/resumes/{pending['resume_id']}/docx")
    ).status_code == 422


def test_custom_lists_compact_spacing_and_missing_contacts_are_native():
    data = {
        "personalInfo": {"name": "仅姓名"},
        "customSections": {
            "talks": {
                "sectionType": "itemList",
                "items": [
                    {
                        "title": "分享会",
                        "subtitle": "参与者",
                        "description": ["介绍调研方法。"],
                    }
                ],
            },
            "tools": {"sectionType": "stringList", "strings": ["SQL", "Excel"]},
        },
        "sectionMeta": [
            {
                "id": "talks",
                "key": "talks",
                "displayName": "分享交流",
                "sectionType": "itemList",
                "isDefault": False,
                "order": 1,
            },
            {
                "id": "tools",
                "key": "tools",
                "displayName": "常用工具",
                "sectionType": "stringList",
                "isDefault": False,
                "order": 2,
            },
        ],
    }
    settings = TemplateSettings(pageSize="LETTER", compactMode=True)
    document = Document(BytesIO(render_resume_docx(data, settings)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "分享会" in text and "SQL、Excel" in text
    assert "@" not in text and "电话" not in text and "example.com" not in text
    assert len(document.inline_shapes) == 0
    assert document.sections[0].page_width.mm == pytest.approx(215.9, abs=0.1)
    assert document.styles["Normal"].paragraph_format.line_spacing.pt == pytest.approx(
        11.25 * 1.45 * 0.92, abs=0.03
    )
    assert document.styles[
        "Heading 1"
    ].paragraph_format.space_before.pt == pytest.approx(15 * 0.6)
    assert document.styles[
        "List Bullet"
    ].paragraph_format.space_after.pt == pytest.approx(6 * 0.6 * 0.75)
    assert copy.deepcopy(data) == data


def test_cjk_fonts_override_theme_defaults_and_title_has_no_border():
    from docx.oxml.ns import qn

    document = Document(
        BytesIO(
            render_resume_docx(
                {
                    "personalInfo": {"name": "中文姓名"},
                    "summary": "中文正文与 English 混排。",
                }
            )
        )
    )
    for name in ("Normal", "Title", "Heading 1", "List Bullet", "List Number"):
        properties = document.styles[name].element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        assert fonts.get(qn("w:eastAsia")) == "Microsoft YaHei"
        assert fonts.get(qn("w:hint")) == "eastAsia"
        assert not any("theme" in key.lower() for key in fonts.attrib)
        assert properties.find(qn("w:lang")).get(qn("w:eastAsia")) == "zh-CN"
    defaults = (
        document.styles.element.find(qn("w:docDefaults"))
        .find(qn("w:rPrDefault"))
        .find(qn("w:rPr"))
    )
    assert defaults.find(qn("w:rFonts")).get(qn("w:eastAsia")) == "Microsoft YaHei"
    assert document.styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr")) is None
    assert document.paragraphs[0]._p.get_or_add_pPr().find(qn("w:pBdr")) is None
    assert "中文正文" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_cjk_exact_line_height_matches_css_without_clipping_photo():
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn

    data = content()
    data["summary"] = "长中文内容与 English 混排，保留清晰的行距。" * 25
    settings = TemplateSettings(fontSize={"base": 5}, spacing={"lineHeight": 3})
    document = Document(BytesIO(render_resume_docx(data, settings)))
    assert document.sections[0]._sectPr.find(qn("w:docGrid")) is None
    for name in ("Normal", "Title", "Heading 1", "List Bullet", "List Number"):
        style = document.styles[name]
        assert (
            style.element.get_or_add_pPr().find(qn("w:snapToGrid")).get(qn("w:val"))
            == "0"
        )
        assert style.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    paragraph_defaults = (
        document.styles.element.find(qn("w:docDefaults"))
        .find(qn("w:pPrDefault"))
        .find(qn("w:pPr"))
    )
    assert paragraph_defaults.find(qn("w:spacing")).get(qn("w:lineRule")) == "exact"
    assert document.styles["Normal"].paragraph_format.line_spacing.pt == pytest.approx(
        16.2
    )
    assert document.styles["Title"].paragraph_format.line_spacing.pt == pytest.approx(
        32.4
    )
    assert document.styles[
        "Heading 1"
    ].paragraph_format.line_spacing.pt == pytest.approx(19.44, abs=0.03)
    picture = document.tables[0].cell(0, 1).paragraphs[0]
    assert picture.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert document.inline_shapes[0].height.mm == pytest.approx(40)
    assert sum(p.text.count("长中文内容") for p in document.paragraphs) == 25
